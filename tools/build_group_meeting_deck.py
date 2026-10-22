#!/usr/bin/env python3
"""Build the 2026-10-22 group meeting deck from a scan and six report figures.

Usage:
    python tools/build_group_meeting_deck.py --report-dir reports/... \
        --batch-dir main/SPD/results/scan_20261020_114505_599047

Requires python-pptx. Source experiments are read only. Slide 1 and slide 6 use
editable text; the other slides embed the exported figures at their native ratio.
All slides contain Chinese speaker notes, also exported to a plain text script.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


REPORT_DATE = "2026-10-22"
FONT = "Microsoft YaHei"
NAVY = "16324F"
TEAL = "147D80"
INK = "273B4F"
MUTED = "62758A"
PALE = "F1F6FA"
WHITE = "FFFFFF"
FIGURE_NAMES = (
    "01_cooperation.png",
    "02_dynamics.png",
    "03_welfare.png",
    "04_group_tradeoff.png",
    "06_advisor.png",
    "05_spatial.png",
)


def color(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def text(slide, x, y, width, height, value, size=23, fill=INK, bold=False):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(9)
        paragraph.line_spacing = 1.15
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(fill)
        # Specify the East Asian typeface too, so Chinese stays editable and
        # uses a standard Windows Chinese font when opened in PowerPoint.
        properties = run._r.get_or_add_rPr()
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", FONT)
        properties.append(east_asian)
    return box


def rectangle(slide, x, y, width, height, fill):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(WHITE)
    return slide


def header(slide, eyebrow, title):
    rectangle(slide, 0.65, 0.48, 0.11, 0.49, TEAL)
    text(slide, 0.94, 0.49, 14.3, 0.4, eyebrow, 15, TEAL, True)
    text(slide, 0.9, 1.12, 14.2, 0.8, title, 35, NAVY, True)


def footer(slide, page, source):
    rectangle(slide, 0.9, 8.38, 14.2, 0.015, "DCE5ED")
    text(slide, 0.9, 8.53, 13.5, 0.25, source, 10, MUTED)
    text(slide, 14.6, 8.5, 0.5, 0.3, f"{page:02d}", 12, MUTED)


def add_notes(slide, value):
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        raise RuntimeError("This presentation template has no speaker-notes body.")
    frame.text = value


def figure_slide(prs, path, notes):
    slide = new_slide(prs)
    picture = slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width)
    # Preserve the figure aspect ratio. Figures produced for this report are
    # 16:9; fitting also avoids distortion if regenerated with a small margin.
    if picture.height > prs.slide_height:
        picture.height = prs.slide_height
        picture.width = int(picture.height * picture.image.size[0] / picture.image.size[1])
    picture.left = (prs.slide_width - picture.width) // 2
    picture.top = (prs.slide_height - picture.height) // 2
    picture.name = path.stem
    add_notes(slide, notes)
    return slide


def read_inputs(batch_dir):
    config = json.loads((batch_dir / "scan_config.json").read_text(encoding="utf-8-sig"))
    with (batch_dir / "summary.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != config.get("completed_runs") or any(
        row.get("status") != "completed" for row in rows
    ):
        raise ValueError("The scan must contain exactly the recorded completed runs.")
    required = {
        "b", "lambda_prob", "seed", "tail_mean_cooperation_rate",
        "delta_population_cumulative_mean_payoff",
        "delta_eligible_cumulative_mean_payoff",
        "delta_ineligible_cumulative_mean_payoff",
        "tail_start_round", "tail_end_round",
    }
    if not rows or not required.issubset(rows[0]):
        raise ValueError("summary.csv is missing fields required for this presentation.")
    # Narrative statements concern this particular single-seed scan, rather
    # than pretending this is a generic statistical report generator.
    if {int(row["seed"]) for row in rows} != {42} or len(rows) != 28:
        raise ValueError("This deck narrative expects 28 parameter settings with seed 42.")
    if {(int(row["tail_start_round"]), int(row["tail_end_round"])) for row in rows} != {(901, 1000)}:
        raise ValueError("This deck narrative expects the tail window 901–1000.")
    return config, rows


def find_run(rows, b, probability):
    matches = [row for row in rows if abs(float(row["b"]) - b) < 1e-10
               and abs(float(row["lambda_prob"]) - probability) < 1e-10]
    if len(matches) != 1:
        raise ValueError(f"Expected one run at b={b}, lambda={probability}.")
    return matches[0]


def fmt_cooperation(row):
    return f"{100 * float(row['tail_mean_cooperation_rate']):.2f}%"


def fmt_delta(row, group):
    return f"{float(row[f'delta_{group}_cumulative_mean_payoff']):+.3f}"


def build_deck(report_dir: Path, batch_dir: Path):
    config, rows = read_inputs(batch_dir)
    figure_dir = report_dir / "figures"
    for name in FIGURE_NAMES:
        if not (figure_dir / name).is_file():
            raise FileNotFoundError(f"Missing report figure: {figure_dir / name}")

    base = find_run(rows, 1.02, 0)
    low = find_run(rows, 1.02, 0.1)
    high = find_run(rows, 1.02, 1)
    case_a = find_run(rows, 1.04, 1)
    case_b = find_run(rows, 1.05, 0.5)
    low_b = find_run(rows, 1.01, 0.5)
    if any(float(row["tail_mean_cooperation_rate"]) != 0
           for row in rows if float(row["b"]) >= 1.04):
        raise ValueError("The expected zero-cooperation observation at b >= 1.04 changed.")

    fixed = config["fixed_parameters"]
    scan = config["scan_values"]
    size = int(fixed["size"])
    access = float(scan["llm_access_fraction"][0])
    steps = int(config["steps"])
    b_values = "、".join(f"{value:g}" for value in scan["b"])
    lambda_values = " / ".join(f"{value:g}" for value in scan["lambda_prob"])
    title = "LLM 建议如何影响空间合作与群体收益？"
    source_label = f"{REPORT_DATE}  ·  来源：{batch_dir.name}  ·  单种子探索性结果"

    notes = [
        (
            "今天汇报最新一批已经完成的扫描。问题是，少数个体有资格咨询 LLM 后，"
            "咨询概率如何改变合作演化，以及收益在资格组和无资格组之间怎样变化。"
            f"设置是 {size}×{size} 周期格点、四近邻相互作用、同步更新，演化 {steps} 步。"
            f"初始合作概率为 {float(fixed['initial_cooperation']):.0%}，固定 {access:.0%} 的位置有咨询资格，"
            f"Fermi 噪声参数 k={float(scan['k'][0]):g}。扫描 b={b_values}，λ={lambda_values}。"
            "λ 是资格位置在更新时使用 LLM 建议的概率，λ=0 是纯 Fermi。"
            "共 28 组参数，每组只有 seed=42 的一次运行；格点和时间步不能当作独立重复。"
            "配对对照有同初态与同资格位置，但干预后随机流消费可能不同。"
            "本次展示描述性观察，不做统计显著性或稳态判断。"
        ),
        (
            "先看最后 100 步，即第 901 到 1000 步的平均合作率。"
            f"b=1.02 时，纯 Fermi 为 {fmt_cooperation(base)}，λ=0.1 为 {fmt_cooperation(low)}，"
            f"λ=1 为 {fmt_cooperation(high)}。同一个 b 下，不同咨询概率可能使合作率向不同方向变化。"
            f"b=1.01、λ=0.5 的末段合作率为 {fmt_cooperation(low_b)}。"
            "在本批扫描的 b≥1.04 条件下，所有末段合作率都是 0。"
            "这里的末段平均不是长期稳态，也不代表已经完成跨种子的验证。"
        ),
        (
            "这一页把所有 b 的完整合作轨迹放在一起，检查末段均值所省略的过程。"
            "重点看合作下降的时间、是否存在暂时维持以及不同 λ 的路径差异。"
            "即使最终或末段合作率同样为 0，前期合作持续的时间仍然不同，"
            "因此累计收益也可能不同。轨迹来自单个种子，不能把每一步视为独立样本。"
        ),
        (
            "收益图的每一格都是实验减去配对纯 Fermi 对照的人均累计收益差，累计第 1 到 1000 步。"
            "三组分别是全体、初始化时固定的资格位置和无资格位置，分组不取决于实际咨询次数。"
            "正值表示该组相同位置的人均累计收益高于对照，负值表示低于对照。"
            "配对差包含邻居传播和整个演化路径改变的作用，不能解释成一次咨询的直接收益。"
            "这一页应与合作率页共同解读：末段合作和全程累计收益是不同指标。"
        ),
        (
            "这里举两个方向相反的组间收益案例。"
            f"b=1.04、λ=1 时，资格组差值 {fmt_delta(case_a, 'eligible')}，"
            f"无资格组 {fmt_delta(case_a, 'ineligible')}，全体 {fmt_delta(case_a, 'population')}。"
            f"b=1.05、λ=0.5 时，资格组 {fmt_delta(case_b, 'eligible')}，"
            f"无资格组 {fmt_delta(case_b, 'ineligible')}，全体 {fmt_delta(case_b, 'population')}。"
            "资格组仅占 20%，所以资格组小幅获益可以与全体收益下降同时出现。"
            "这只是两个单次运行中的描述性案例，尤其小幅差值还需要多种子检验。"
            "同时，差值不等于资格带来的纯直接效应，仍包含系统传播。"
        ),
        (
            "目前有三个谨慎结论。第一，LLM 建议的效果取决于 b 和咨询概率，"
            "当前数据不支持建议越多越好。第二，只看末段合作率会漏掉演化过程和累计收益的差异。"
            "第三，资格组受益并不自动等于全体受益，后续应继续分组评估。"
            "下一步先做多 seed 配对重复：先计算每个 seed 的实验减对照，再在 seed 之间汇总均值与区间。"
            "同时独立重复生成建议策略表，区分演化随机性与建议表随机性，"
            "并增加每个局面的独立回复数。当前不报告统计显著性，不把末 100 步称为稳态。"
        ),
        (
            "附录展示本次使用的建议策略表。建议来源为 gpt-5.5，输入为 3×3 二元局面。"
            "全部 512 个局面各生成 2 条回复，1024 条回复全部成功。"
            "108 个局面的两条回复都是合作，340 个都是背叛，64 个不一致。"
            "模拟中对当前局面对应的两条固定回复有放回抽样，重复使用这张表。"
            "合作建议共 280 条，占 27.34375%；本人原为合作时合作建议占 48.828125%，"
            "本人原为背叛时占 5.859375%。这些比例按表内局面计算，不等于模拟中实际访问的局面分布。"
            "每局面只有 2 条回复，不能据此精确估计模型在该局面下的选择概率。"
            "后续应独立重新生成策略表，并增加每个局面的回复数。"
        ),
        (
            "最后用最终空间状态辅助解释合作的空间分布。"
            "每张格点图只是第 1000 步的一个快照，可以观察合作团簇或全 D 状态，"
            "但单张快照本身不能证明形成机制，也不能证明已经达到长期稳态。"
            "解释时应结合前面的完整动态轨迹、末段均值和配对累计收益。"
        ),
    ]

    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    prs.core_properties.title = title
    prs.core_properties.subject = "2026-10-22 组会：最新空间囚徒困境参数扫描"
    prs.core_properties.author = "LLM_Advice"
    prs.core_properties.keywords = "空间合作,LLM,Fermi,组会,单种子探索"

    slide = new_slide(prs)
    header(slide, "组会汇报  /  最新参数扫描  /  2026-10-22", title)
    text(slide, 0.94, 2.12, 14.1, 0.55,
         "关注三件事：合作能否维持、累计收益如何变化、收益归于谁。", 24, MUTED)
    rectangle(slide, 0.9, 3.00, 8.75, 3.92, PALE)
    text(slide, 1.25, 3.25, 8.0, 0.5, "实验设置", 24, NAVY, True)
    text(slide, 1.25, 3.95, 8.0, 2.8,
         f"{size} × {size} 格点｜周期边界｜四近邻｜同步更新\n"
         f"演化 {steps:,} 步；初始合作概率 {float(fixed['initial_cooperation']):.0%}；k = {float(scan['k'][0]):g}\n"
         f"固定 {access:.0%} 位置有咨询资格（{round(size * size * access)} / {size * size}）\n"
         f"λ = {lambda_values}；λ = 0 为纯 Fermi\n"
         f"b = {b_values}", 22)
    rectangle(slide, 10.00, 3.00, 5.10, 3.92, NAVY)
    text(slide, 10.40, 3.28, 4.25, 0.7, "28 组参数", 32, WHITE, True)
    text(slide, 10.40, 4.22, 4.25, 2.2,
         "每组一次运行：seed = 42\n配对纯 Fermi 对照\n末段均值：第 901–1000 步", 22, WHITE)
    text(slide, 0.95, 7.32, 14.1, 0.65,
         "本轮定位：探索性观察；不做显著性判断，不将末段均值称为稳态。", 21, TEAL, True)
    footer(slide, 1, source_label)
    add_notes(slide, notes[0])

    for index, name in enumerate(FIGURE_NAMES[:4], start=1):
        figure_slide(prs, figure_dir / name, notes[index])

    slide = new_slide(prs)
    header(slide, "当前结论与下一步", "三个观察，仍需多种子与策略表重复验证")
    conclusions = (
        ("01", "效果依赖 b 与咨询概率",
         f"b = 1.02：Fermi {fmt_cooperation(base)} → λ = 0.1 时 {fmt_cooperation(low)}；"
         f"λ = 1 时 {fmt_cooperation(high)}。"),
        ("02", "末段合作率不能概括全程收益",
         "本批 b ≥ 1.04 的末 100 步合作率均为 0；不同条件的累计收益差仍有正有负。"),
        ("03", "资格组受益不等于全体受益",
         f"b = 1.04，λ = 1：资格组 {fmt_delta(case_a, 'eligible')}，"
         f"无资格组 {fmt_delta(case_a, 'ineligible')}，全体 {fmt_delta(case_a, 'population')}。"),
    )
    for index, (number, heading, body) in enumerate(conclusions):
        y = 2.22 + index * 1.46
        rectangle(slide, 0.9, y, 14.2, 1.21, PALE)
        text(slide, 1.18, y + 0.22, 0.8, 0.65, number, 29, TEAL, True)
        text(slide, 2.05, y + 0.15, 12.5, 0.42, heading, 23, NAVY, True)
        text(slide, 2.05, y + 0.69, 12.5, 0.36, body, 19, INK)
    text(slide, 0.98, 6.96, 14.0, 0.55,
         "下一步：多 seed 配对重复 → 在 seed 间汇总区间；独立重复生成策略表。", 23, TEAL, True)
    text(slide, 0.98, 7.65, 14.0, 0.42,
         "解释边界：配对差包含邻居传播；格点与时间步不作为独立重复。", 18, MUTED)
    footer(slide, 6, source_label)
    add_notes(slide, notes[5])
    figure_slide(prs, figure_dir / FIGURE_NAMES[4], notes[6])
    figure_slide(prs, figure_dir / FIGURE_NAMES[5], notes[7])

    report_dir.mkdir(parents=True, exist_ok=True)
    destination = report_dir / "group_meeting_20261022.pptx"
    prs.save(destination)
    titles = (
        "实验问题与设置", "末 100 步合作率", "完整合作演化轨迹",
        "相对配对 Fermi 的人均累计收益差", "组间收益分化案例",
        "三个审慎结论与下一步", "附录：建议策略表", "附录：最终空间状态",
    )
    script = report_dir / "组会讲稿.txt"
    blocks = [f"组会讲稿｜{REPORT_DATE}\n数据来源：{batch_dir.name}\n建议：前 6 页主讲，后 2 页按需补充。"]
    for index, (slide_title, note) in enumerate(zip(titles, notes), start=1):
        blocks.append(f"第 {index} 页｜{slide_title}\n{note}")
    script.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
    return destination, script


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report-dir", type=Path, required=True)
    parser.add_argument("--batch-dir", type=Path, required=True,
                        help="Scan directory containing summary.csv and scan_config.json")
    args = parser.parse_args()
    for path in build_deck(args.report_dir.resolve(), args.batch_dir.resolve()):
        print(path)


if __name__ == "__main__":
    main()
